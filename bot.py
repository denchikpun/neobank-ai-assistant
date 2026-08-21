import os
import re
import json
import asyncio
import logging
import tempfile
import requests
from datetime import datetime

from groq import Groq
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, BotCommand
from telegram.ext import (
    Application, CommandHandler, MessageHandler, CallbackQueryHandler,
    ContextTypes, filters, ConversationHandler
)

logging.basicConfig(format="%(asctime)s | %(levelname)s | %(message)s", level=logging.INFO)
logger = logging.getLogger(__name__)

TELEGRAM_TOKEN  = os.environ["TELEGRAM_BOT_TOKEN"]
GROQ_KEY        = os.environ["GROQ_API_KEY"]
APPS_SCRIPT_URL = os.environ.get("APPS_SCRIPT_URL", "")
SCRIPT_TOKEN    = os.environ.get("SCRIPT_TOKEN", "TRUSTME_SECRET_2025")
TAVILY_KEY      = os.environ.get("TAVILY_API_KEY", "")
ALLOWED_USERNAMES = set(
    u.strip().lstrip("@").lower()
    for u in os.environ.get("ALLOWED_USERNAMES", "denya951,daukaz,kompassito").split(",")
    if u.strip()
)
# superadmin can delete any folder; admins can delete only folders they created
SUPERADMIN = os.environ.get("SUPERADMIN_USERNAME", "denya951").strip().lstrip("@").lower()

# Google Drive API (migration) — optional; when set, enables direct Drive access
GOOGLE_SA_JSON = os.environ.get("GOOGLE_SA_JSON", "")
SHARED_DRIVE_FOLDER_ID = os.environ.get("SHARED_DRIVE_FOLDER_ID", "")

groq_client = Groq(api_key=GROQ_KEY)


def drive_api_upload(file_path, file_name, folder_id, mime_type):
    """Upload a file directly to Drive via the API (no base64/Apps Script).
    Used for large files. Returns (file_id, web_url, error)."""
    service, err = get_drive_service()
    if err:
        return None, None, err
    try:
        from googleapiclient.http import MediaFileUpload
        media = MediaFileUpload(file_path, mimetype=mime_type or "application/octet-stream",
                                resumable=True)
        created = service.files().create(
            body={"name": file_name, "parents": [folder_id]},
            media_body=media,
            fields="id, webViewLink",
            supportsAllDrives=True,
        ).execute()
        return created.get("id"), created.get("webViewLink"), None
    except Exception as e:
        return None, None, f"{type(e).__name__}: {str(e)[:200]}"


def resolve_folder_id(folder_path):
    """Ask Apps Script for the Drive ID of a folder path."""
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "resolve_folder", "folder": folder_path
        }, timeout=15)
        r = resp.json()
        return (r.get("folder_id"), None) if r.get("ok") else (None, r.get("error"))
    except Exception as e:
        return None, str(e)


def get_drive_service():
    """Build a Google Drive API client from the Service Account JSON.
    Returns (service, error). service is None if not configured or failed."""
    if not GOOGLE_SA_JSON:
        return None, "GOOGLE_SA_JSON not set"
    try:
        import json as _json
        from google.oauth2 import service_account
        from googleapiclient.discovery import build
        info = _json.loads(GOOGLE_SA_JSON)
        creds = service_account.Credentials.from_service_account_info(
            info, scopes=["https://www.googleapis.com/auth/drive"]
        )
        service = build("drive", "v3", credentials=creds, cache_discovery=False)
        return service, None
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


def drive_api_selftest():
    """Read+write test against SHARED_DRIVE_FOLDER_ID. Returns a result string."""
    service, err = get_drive_service()
    if err:
        return f"❌ Drive API not ready: {err}"
    if not SHARED_DRIVE_FOLDER_ID:
        return "❌ SHARED_DRIVE_FOLDER_ID not set"

    # identify what this ID actually is
    info_line = ""
    is_shared_drive_root = False
    try:
        meta = service.files().get(
            fileId=SHARED_DRIVE_FOLDER_ID,
            fields="id, name, mimeType, driveId, parents, capabilities(canAddChildren)",
            supportsAllDrives=True,
        ).execute()
        is_shared_drive_root = (meta.get("mimeType") == "application/vnd.google-apps.folder"
                                and not meta.get("parents")
                                and meta.get("driveId") == SHARED_DRIVE_FOLDER_ID)
        info_line = (f"ID is: {meta.get('name')} | type={meta.get('mimeType')} | "
                     f"driveId={meta.get('driveId','none')} | "
                     f"canAddChildren={meta.get('capabilities',{}).get('canAddChildren')}")
    except Exception:
        # a Shared Drive root often isn't returned by files().get — try drives().get
        try:
            d = service.drives().get(driveId=SHARED_DRIVE_FOLDER_ID).execute()
            is_shared_drive_root = True
            info_line = f"ID is a SHARED DRIVE root: {d.get('name')}"
        except Exception as e2:
            info_line = f"could not identify ID: {str(e2)[:150]}"

    # read test
    try:
        resp = service.files().list(
            q=f"'{SHARED_DRIVE_FOLDER_ID}' in parents and trashed=false",
            fields="files(id, name)", pageSize=5,
            supportsAllDrives=True, includeItemsFromAllDrives=True,
        ).execute()
        n = len(resp.get("files", []))
    except Exception as e:
        return f"❌ READ failed: {type(e).__name__}: {str(e)[:200]}\n\n{info_line}"

    # write test
    try:
        import io
        from googleapiclient.http import MediaIoBaseUpload
        content = io.BytesIO(b"TrustMe Drive API test - write ok")
        media = MediaIoBaseUpload(content, mimetype="text/plain", resumable=False)
        body = {"name": "trustme_drive_api_test.txt", "parents": [SHARED_DRIVE_FOLDER_ID]}
        created = service.files().create(
            body=body, media_body=media, fields="id, webViewLink",
            supportsAllDrives=True,
        ).execute()
        return (f"✅ Drive API works!\n"
                f"Read: saw {n} item(s)\n"
                f"Write: created test file\n"
                f"{created.get('webViewLink')}\n\n{info_line}")
    except Exception as e:
        # list the Shared Drives the Service Account is actually a member of
        member_of = ""
        try:
            drives = service.drives().list(pageSize=20, fields="drives(id,name)").execute()
            ds = drives.get("drives", [])
            if ds:
                member_of = "\n\nService Account is a member of these Shared Drives:\n" + \
                    "\n".join([f"• {d['name']} (id: {d['id']})" for d in ds])
            else:
                member_of = ("\n\n⚠️ Service Account is a member of NO Shared Drives. "
                             "It was shared the folder directly, not added to the Drive. "
                             "Add it via the Shared Drive's Manage Members.")
        except Exception as e2:
            member_of = f"\n\n(could not list drives: {str(e2)[:120]})"
        hint = ""
        if is_shared_drive_root:
            hint = "\n\n💡 This is the Drive ROOT — use a folder inside it instead."
        return (f"⚠️ READ ok but WRITE failed: {type(e).__name__}: {str(e)[:200]}\n\n"
                f"{info_line}{hint}{member_of}")
MODEL = "llama-3.3-70b-versatile"

WAITING_FOCUS = 1

FOLDERS = {
    "Knowledge/Product":       "Product specs, PRD, roadmap, features, UX",
    "Knowledge/Sharia":        "Sharia standards, fatwa, AAOIFI, IFSB, halal screening",
    "Knowledge/Market":        "Competitors, market research, fintech trends, analysis",
    "Knowledge/Strategy":      "Pitch decks, OKR, fundraising, vision, investors",
    "Knowledge/User_Insights": "User feedback, research, interviews, surveys",
    "Operations/Decisions":    "Team decisions, choices, rationale",
    "Operations/Meetings":     "Meeting notes, agendas, minutes",
    "Operations/Processes":    "SOPs, workflows, instructions, guides",
    "Operations/Agents":       "AI agent instructions and configurations",
    "Operations/HR":           "Team structure, roles, onboarding",
    "Company/Product_Tech":    "Product and technology department docs",
    "Company/Compliance":      "Compliance and Sharia department docs",
    "Company/Banking_Ops":     "Banking operations department docs",
    "Company/Investments":     "Investments and Halal ETF department docs",
    "Company/Marketing":       "Marketing and growth department docs",
    "Company/Finance_IR":      "Finance and investor relations docs",
    "Company/HR":              "HR and org structure docs",
}
FOLDERS_TEXT = "\n".join([f"- {k}: {v}" for k, v in FOLDERS.items()])


# cache for the folder list — scanning Drive is slow, and the structure
# rarely changes, so we hold the result for a short window
_folder_cache = {"folders": None, "ts": 0.0}
FOLDER_CACHE_TTL = 300  # seconds (5 min); cleared immediately on folder create/delete


def invalidate_folder_cache():
    _folder_cache["folders"] = None
    _folder_cache["ts"] = 0.0


def get_live_folders(force=False):
    """Return the current folder list: hardcoded seed + any created in Drive.
    Cached in memory for FOLDER_CACHE_TTL seconds to avoid repeated slow scans.
    Falls back to the hardcoded list if the scan fails."""
    import time as _time
    base = list(FOLDERS.keys())
    if not APPS_SCRIPT_URL:
        return base
    # serve from cache if fresh
    if not force and _folder_cache["folders"] is not None:
        if (_time.time() - _folder_cache["ts"]) < FOLDER_CACHE_TTL:
            return _folder_cache["folders"]
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "list_folders"
        }, timeout=20)
        result = resp.json()
        if result.get("ok"):
            scanned = list(result.get("folders", {}).keys())
            merged = base + [f for f in scanned if f not in base]
            merged = merged or base
            _folder_cache["folders"] = merged
            _folder_cache["ts"] = _time.time()
            return merged
    except Exception as e:
        logger.info(f"list_folders failed, using hardcoded: {e}")
    # on failure, serve stale cache if we have it, else the seed
    return _folder_cache["folders"] or base


def is_superadmin(update):
    u = getattr(update, "effective_user", None)
    return bool(u) and (u.username or "").lower() == SUPERADMIN


def register_user(update):
    """Remember this user's chat_id in Drive so we can notify them later.
    Only registers allowed admins (the notification audience)."""
    if not APPS_SCRIPT_URL:
        return
    u = getattr(update, "effective_user", None)
    chat = getattr(update, "effective_chat", None)
    if not u or not chat:
        return
    uname = (u.username or "").lower()
    if uname not in ALLOWED_USERNAMES:
        return
    try:
        requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "register_user",
            "chat_id": chat.id, "username": uname,
        }, timeout=10)
    except Exception as e:
        logger.info(f"register_user failed: {e}")


def _get_user(obj):
    """Works with both Update (effective_user) and CallbackQuery (from_user)."""
    return getattr(obj, "effective_user", None) or getattr(obj, "from_user", None)


def _display_name(update):
    u = _get_user(update)
    if not u:
        return "Someone"
    return u.first_name or ("@" + u.username if u.username else "Someone")


async def notify_admins(context, update, text):
    """Send a notification to all registered admins except the author.
    Returns a short diagnostic string (sent count / errors)."""
    if not APPS_SCRIPT_URL:
        return "no apps script url"
    author_user = _get_user(update)
    author = (getattr(author_user, "username", "") or "").lower()
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "list_users"
        }, timeout=10)
        users = resp.json().get("users", {})
    except Exception as e:
        logger.info(f"notify list_users failed: {e}")
        return f"list_users failed: {e}"
    sent, errors = 0, []
    for uname, chat_id in users.items():
        if uname == author:
            continue  # don't notify the person who did the action
        try:
            await context.bot.send_message(chat_id=int(chat_id), text=text,
                                           disable_web_page_preview=True)
            sent += 1
        except Exception as e:
            errors.append(f"{uname}: {str(e)[:80]}")
            logger.info(f"notify to {uname} ({chat_id}) failed: {e}")
    diag = f"notified {sent}"
    if errors:
        diag += " | errors: " + "; ".join(errors)
    logger.info(f"notify_admins: {diag}")
    return diag


def folders_for_prompt():
    """Folder list text for AI classification — includes live sub-folders.
    Falls back to the hardcoded descriptions when scan is unavailable."""
    live = get_live_folders()
    lines = []
    for path in live:
        desc = FOLDERS.get(path, "")
        lines.append(f"- {path}: {desc}" if desc else f"- {path}")
    return "\n".join(lines)


def build_folder_menu(prefix, allow_new, newfolder_cb):
    """Top-level folder menu. Folders that have sub-folders get an 'expand'
    button (▸) leading to their sub-folders; leaf folders select directly.
    prefix is 'folder' (analysis) or 'rawfolder' (raw save)."""
    live = get_live_folders()
    tops = [f for f in live if f.count("/") == 1]  # e.g. Company/Banking_Ops
    rows = []
    for top in tops:
        has_subs = any(f.startswith(top + "/") for f in live)
        if has_subs:
            rows.append([InlineKeyboardButton(f"{top}  ▸", callback_data=f"expand_|{prefix}|{top}")])
        else:
            rows.append([InlineKeyboardButton(top, callback_data=f"{prefix}_{top}")])
    if allow_new:
        rows.append([InlineKeyboardButton("➕ New folder", callback_data=newfolder_cb)])
    return rows


def create_folder_in_drive(parent_path, name, created_by):
    """Ask Apps Script to create a folder. Returns (path, error)."""
    if not APPS_SCRIPT_URL:
        return None, "Apps Script URL not configured"
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "create_folder",
            "parent_path": parent_path, "name": name, "created_by": created_by,
        }, timeout=30)
        result = resp.json()
        if result.get("ok"):
            return result.get("path"), None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


def delete_folder_in_drive(path, requested_by, superadmin):
    """Ask Apps Script to archive a folder. Returns (result, error)."""
    if not APPS_SCRIPT_URL:
        return None, "Apps Script URL not configured"
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "delete_folder",
            "path": path, "requested_by": requested_by, "is_superadmin": superadmin,
        }, timeout=30)
        result = resp.json()
        if result.get("ok"):
            return result, None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


def is_allowed(update):
    """Access is granted only to Telegram usernames in the allowlist."""
    user = getattr(update, "effective_user", None)
    if user is None:
        return False
    uname = (user.username or "").lower()
    # empty allowlist = open (safety fallback); normally it's populated
    if not ALLOWED_USERNAMES:
        return True
    return uname in ALLOWED_USERNAMES


async def deny(update):
    """Tell an unauthorized user they don't have access."""
    try:
        await update.message.reply_text(
            "⛔️ Access denied. You are not authorized to use this bot.\n"
            "Contact an administrator to request access."
        )
    except Exception:
        pass


VIDEO_HOSTS = ("youtube.com", "youtu.be", "vimeo.com", "tiktok.com",
               "instagram.com", "facebook.com/watch", "twitch.tv")


def is_video_url(url):
    u = url.lower()
    return any(h in u for h in VIDEO_HOSTS)


def fetch_url_content(url):
    """Return extracted text, or None if the page could not be read."""
    try:
        headers = {"User-Agent": "Mozilla/5.0 (compatible; TrustMeBot/1.0)"}
        resp = requests.get(url, headers=headers, timeout=15)
        resp.raise_for_status()
        html = resp.text

        # 1) Try trafilatura — extracts main article, strips nav/footer/menus
        try:
            import trafilatura
            extracted = trafilatura.extract(
                html,
                include_comments=False,
                include_tables=False,
                no_fallback=False,
            )
            if extracted and len(extracted.strip()) > 200:
                return extracted.strip()[:40000]
        except Exception as e:
            logger.info(f"trafilatura failed, falling back: {e}")

        # 2) Fallback — crude tag strip (last resort)
        text = re.sub(r"<script[^>]*>.*?</script>", "", html, flags=re.DOTALL)
        text = re.sub(r"<style[^>]*>.*?</style>",   "", text, flags=re.DOTALL)
        text = re.sub(r"<nav[^>]*>.*?</nav>",       "", text, flags=re.DOTALL)
        text = re.sub(r"<footer[^>]*>.*?</footer>", "", text, flags=re.DOTALL)
        text = re.sub(r"<header[^>]*>.*?</header>", "", text, flags=re.DOTALL)
        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"\s+",     " ", text).strip()
        # only accept if we actually got meaningful text
        if len(text) > 200:
            return text[:40000]
        return None
    except Exception as e:
        logger.info(f"fetch_url_content failed for {url}: {e}")
        return None


def extract_file_content(tmp_path, filename):
    fname = filename.lower()
    try:
        if fname.endswith(".pdf"):
            try:
                import fitz
                doc = fitz.open(tmp_path)
                return "\n".join([p.get_text() for p in doc])[:40000]
            except ImportError:
                pass
        if fname.endswith(".docx"):
            try:
                from docx import Document as DocxDoc
                doc = DocxDoc(tmp_path)
                return "\n".join([p.text for p in doc.paragraphs])[:40000]
            except ImportError:
                pass
        # Modern Excel formats (openpyxl): .xlsx .xlsm .xltx .xltm
        if fname.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
            try:
                import openpyxl
            except ImportError as e:
                extract_file_content.last_error = f"openpyxl not installed: {e}"
                logger.info(extract_file_content.last_error)
                return None
            try:
                wb = openpyxl.load_workbook(tmp_path, read_only=True, data_only=True)
                out = []
                for ws in wb.worksheets:
                    out.append(f"=== Sheet: {ws.title} ===")
                    for row in ws.iter_rows(values_only=True):
                        cells = [str(c) for c in row if c is not None]
                        if cells:
                            out.append(" | ".join(cells))
                    if sum(len(x) for x in out) > 40000:
                        break
                wb.close()
                text = "\n".join(out).strip()
                if text:
                    return text[:40000]
                extract_file_content.last_error = "spreadsheet opened but no readable cells found"
                return None
            except Exception as e:
                extract_file_content.last_error = f"openpyxl read error: {type(e).__name__}: {e}"
                logger.info(extract_file_content.last_error)
                return None
        # Legacy Excel 97-2003 (.xls) via xlrd
        if fname.endswith(".xls"):
            try:
                import xlrd
            except ImportError as e:
                extract_file_content.last_error = f"xlrd not installed: {e}"
                logger.info(extract_file_content.last_error)
                return None
            try:
                wb = xlrd.open_workbook(tmp_path)
                out = []
                for sh in wb.sheets():
                    out.append(f"=== Sheet: {sh.name} ===")
                    for r in range(sh.nrows):
                        cells = [str(sh.cell_value(r, c)) for c in range(sh.ncols)
                                 if sh.cell_value(r, c) != ""]
                        if cells:
                            out.append(" | ".join(cells))
                    if sum(len(x) for x in out) > 40000:
                        break
                text = "\n".join(out).strip()
                if text:
                    return text[:40000]
                extract_file_content.last_error = "xls opened but no readable cells found"
                return None
            except Exception as e:
                extract_file_content.last_error = f"xlrd read error: {type(e).__name__}: {e}"
                logger.info(extract_file_content.last_error)
                return None
        # OpenDocument Spreadsheet (.ods) via odfpy
        if fname.endswith(".ods"):
            try:
                from odf.opendocument import load as odf_load
                from odf.table import Table, TableRow, TableCell
                from odf.text import P
            except ImportError as e:
                extract_file_content.last_error = f"odfpy not installed: {e}"
                logger.info(extract_file_content.last_error)
                return None
            try:
                doc = odf_load(tmp_path)
                out = []
                for table in doc.spreadsheet.getElementsByType(Table):
                    out.append(f"=== Sheet: {table.getAttribute('name')} ===")
                    for row in table.getElementsByType(TableRow):
                        cells = []
                        for cell in row.getElementsByType(TableCell):
                            txt = "".join(str(p) for p in cell.getElementsByType(P))
                            if txt:
                                cells.append(txt)
                        if cells:
                            out.append(" | ".join(cells))
                    if sum(len(x) for x in out) > 40000:
                        break
                text = "\n".join(out).strip()
                if text:
                    return text[:40000]
                extract_file_content.last_error = "ods opened but no readable cells found"
                return None
            except Exception as e:
                extract_file_content.last_error = f"ods read error: {type(e).__name__}: {e}"
                logger.info(extract_file_content.last_error)
                return None
        # Rich Text Format (.rtf)
        if fname.endswith(".rtf"):
            try:
                from striprtf.striprtf import rtf_to_text
            except ImportError as e:
                extract_file_content.last_error = f"striprtf not installed: {e}"
                logger.info(extract_file_content.last_error)
                return None
            try:
                with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
                    raw = f.read()
                text = rtf_to_text(raw).strip()
                if text:
                    return text[:40000]
                extract_file_content.last_error = "rtf parsed but empty"
                return None
            except Exception as e:
                extract_file_content.last_error = f"rtf read error: {type(e).__name__}: {e}"
                logger.info(extract_file_content.last_error)
                return None
        # Plain delimited text
        if fname.endswith((".csv", ".tsv", ".txt", ".md", ".json")):
            with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read(40000)
        with open(tmp_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read(40000)
    except Exception as e:
        extract_file_content.last_error = f"{type(e).__name__}: {e}"
        logger.info(f"extract_file_content failed for {filename}: {e}")
        return None

extract_file_content.last_error = ""


def looks_like_garbage(text):
    """True if text looks like markup/binary junk rather than readable content.
    Prevents the AI from hallucinating a confident summary of unreadable input."""
    if not text:
        return True
    sample = text[:2000]
    # RTF / control-word markup signatures
    if sample.lstrip().startswith(("{\\rtf", "{\\*", "%PDF", "PK\x03\x04")):
        return True
    # ratio of "structural" chars typical of markup/binary (pipe excluded — tables use it)
    junk = sum(sample.count(c) for c in "{}\\<>")
    if junk > len(sample) * 0.08:
        return True
    # ratio of readable letters (any alphabet) must be reasonable
    letters = sum(1 for c in sample if c.isalpha() or c.isspace())
    if letters < len(sample) * 0.5:
        return True
    return False


def process_with_ai(content, source, focus):
    pages  = max(1, len(content.split()) // 250)
    n_tags = "15-20" if pages >= 100 else "10-15" if pages >= 50 else "5-10" if pages >= 10 else "3-5"
    today  = datetime.now().strftime("%Y-%m-%d")
    folder_list = folders_for_prompt()

    prompt = f"""You are the NeoBank Knowledge Agent. Process this material and return ONLY valid JSON — no markdown, no explanation.

LANGUAGE RULE: The input material may be in ANY language (Russian, English, Arabic, etc.). You must UNDERSTAND it in its original language, but ALWAYS write every output field (title, summary, key_insights, hashtags) in ENGLISH. Never output non-English text.

SOURCE: {source}
USER FOCUS: {focus if focus else "General — extract what is most relevant for TrustMe Islamic Digital Bank"}
DATE: {today}

MATERIAL:
{content[:30000]}

Available folders:
{FOLDERS_TEXT}

Return ONLY this JSON:
{{
  "title": "concise English title",
  "folder": "one folder path from the list",
  "summary": "3-5 sentences focused on TrustMe relevance",
  "key_insights": ["insight 1", "insight 2", "insight 3"],
  "hashtags": [{n_tags} strings like "#topic_sharia" "#type_article" "#dept_product" "#key_halal_etf" "#lang_en"],
  "importance": <1-10>,
  "importance_reason": "one sentence",
  "credibility": <1-10>,
  "credibility_reason": "one sentence",
  "version": "V1.0",
  "date": "{today}"
}}

Scoring rules — follow these scales exactly (from the agent's operating rules):

IMPORTANCE (1-10):
- 9-10 Critical: Sharia standards, company strategy, foundational decisions
- 7-8  High: key research, product decisions, competitive analysis
- 5-6  Medium: industry trends, useful background, supporting data
- 3-4  Low: general context, tangential information
- 1-2  Minimal: archival reference only

CREDIBILITY (1-10) — judge by SOURCE TYPE:
- 9-10 Official/Regulatory: AAOIFI, IFSB, regulatory docs, internal team decisions
- 7-8  Academic: peer-reviewed articles, analytical agency reports
- 5-7  Media/Publications: business media, authoritative industry blogs
- 4-6  Interview/Podcast: single expert opinion (reduce accordingly)
- 1-3  Social/Forums: Twitter, LinkedIn, forums (requires verification)

Other rules: 3-7 key insights; every output field in English."""

    response = groq_client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.4,
        response_format={"type": "json_object"},
    )
    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```json\s*", "", raw)
    raw = re.sub(r"^```\s*",     "", raw)
    raw = re.sub(r"\s*```$",     "", raw)
    return json.loads(raw)


def save_to_drive(data, source, content):
    """Send document to Apps Script which saves it to Drive."""
    if not APPS_SCRIPT_URL:
        return None, "Apps Script URL not configured"

    payload = {
        "token":       SCRIPT_TOKEN,
        "action":      "create_doc",
        "title":       data.get("title", "Untitled"),
        "folder":      data.get("folder", "Knowledge/Market"),
        "content":     f"SUMMARY:\n{data.get('summary','')}\n\nKEY INSIGHTS:\n" +
                       "\n".join([f"{i+1}. {t}" for i,t in enumerate(data.get('key_insights',[]))]) +
                       f"\n\nSOURCE CONTENT (cleaned excerpt):\n{content[:6000]}",
        "source":      source,
        "importance":  data.get("importance", 5),
        "credibility": data.get("credibility", 5),
        "date":        data.get("date", datetime.now().strftime("%Y-%m-%d")),
        "version":     data.get("version", "V1.0"),
        "hashtags":    data.get("hashtags", []),
        "summary":     data.get("summary", ""),
        "key_insights": data.get("key_insights", []),
    }

    try:
        resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=30)
        logger.info(f"Apps Script response [{resp.status_code}]: {resp.text[:500]}")
        result = resp.json()
        if result.get("ok"):
            return result.get("file_url"), result.get("file_id"), None
        else:
            return None, None, result.get("error", "Unknown error from Apps Script")
    except Exception as e:
        return None, None, str(e)


# formats where the original file must be preserved (structure matters)
STRUCTURAL_EXTS = (".xlsx", ".xlsm", ".xls", ".xltx", ".xltm", ".ods",
                   ".pdf", ".pptx", ".ppt", ".odp")

MIME_BY_EXT = {
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xls":  "application/vnd.ms-excel",
    ".ods":  "application/vnd.oasis.opendocument.spreadsheet",
    ".pdf":  "application/pdf",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt":  "application/vnd.ms-powerpoint",
    ".odp":  "application/vnd.oasis.opendocument.presentation",
    ".ogg":  "audio/ogg",
    ".oga":  "audio/ogg",
    ".mp3":  "audio/mpeg",
    ".wav":  "audio/wav",
    ".jpg":  "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png":  "image/png",
    ".webp": "image/webp",
}

# Telegram Bot API hard limit — bots cannot download files larger than 20 MB
TELEGRAM_DOWNLOAD_LIMIT = 20 * 1024 * 1024
# Apps Script base64 request cap — keep originals under ~30 MB raw
MAX_ORIGINAL_BYTES = 30 * 1024 * 1024  # ceiling for the base64/Apps Script path
# files larger than this use the Drive API directly (no base64) when configured
DRIVE_API_THRESHOLD = 8 * 1024 * 1024  # 8 MB — above this, prefer Drive API
# with Drive API available, allow much larger originals
MAX_DRIVE_API_BYTES = 100 * 1024 * 1024  # 100 MB via direct Drive API


def is_structural_file(filename):
    return (filename or "").lower().endswith(STRUCTURAL_EXTS)


def make_filename(base_title, version="V1.0"):
    """Build a stored name per our rules: Title_Version_YYYY-MM-DD (no extension —
    Apps Script re-adds the original extension). Cleans spaces/punctuation."""
    # strip any existing extension from the incoming name
    base = os.path.splitext(base_title or "Untitled")[0]
    # keep letters/numbers/underscore; turn spaces and separators into underscore
    base = re.sub(r"[^\w\-]+", "_", base, flags=re.UNICODE).strip("_") or "Untitled"
    date = datetime.now().strftime("%Y-%m-%d")
    return f"{base}_{version}_{date}"


def save_file_only_to_drive(folder, title, file_path, file_name, source,
                            importance=0, credibility=0):
    """Save just the original file to Drive (no analysis doc)."""
    if not APPS_SCRIPT_URL:
        return None, None, "Apps Script URL not configured"
    try:
        size = os.path.getsize(file_path)
        ext = os.path.splitext(file_name)[1].lower()
        mime = MIME_BY_EXT.get(ext, "application/octet-stream")

        # large file → Drive API direct upload
        if size > DRIVE_API_THRESHOLD and GOOGLE_SA_JSON:
            if size > MAX_DRIVE_API_BYTES:
                return None, None, f"file too large ({size // (1024*1024)} MB > 100 MB)"
            folder_id, ferr = resolve_folder_id(folder)
            if not folder_id:
                return None, None, f"could not resolve folder: {ferr}"
            stored_name = title if title.lower().endswith(ext) else (title + ext)
            up_id, up_url, uerr = drive_api_upload(file_path, stored_name, folder_id, mime)
            if uerr:
                return None, None, f"Drive API upload failed: {uerr}"
            # log to changelog via Apps Script (best-effort)
            try:
                requests.post(APPS_SCRIPT_URL, json={
                    "token": SCRIPT_TOKEN, "action": "log_raw",
                    "title": title, "folder": folder, "file_url": up_url,
                    "importance": importance, "credibility": credibility,
                }, timeout=30)
            except Exception:
                pass
            return up_url, up_id, None

        import base64
        if size > MAX_ORIGINAL_BYTES:
            return None, None, f"file too large ({size // (1024*1024)} MB > 30 MB)"
        with open(file_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("ascii")
        payload = {
            "token": SCRIPT_TOKEN,
            "action": "save_file_only",
            "title": title,
            "folder": folder,
            "file_b64": b64,
            "file_name": file_name,
            "mime_type": MIME_BY_EXT.get(ext, "application/octet-stream"),
            "date": datetime.now().strftime("%Y-%m-%d"),
            "version": "V1.0",
            "source": source,
            "importance": importance,
            "credibility": credibility,
            "hashtags": ["#raw_unprocessed"],
        }
        resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=300)
        logger.info(f"save_file_only response [{resp.status_code}]: {resp.text[:300]}")
        try:
            result = resp.json()
        except Exception:
            return None, None, f"[{resp.status_code}] non-JSON: {resp.text[:200]}"
        if result.get("ok"):
            url = result.get("file_url")
            if url:
                return url, result.get("file_id"), None
            # ok but no url — surface what came back so we can see the field names
            return None, None, f"saved but no URL in response: {str(result)[:200]}"
        return None, None, result.get("error") or f"unknown; raw: {str(result)[:200]}"
    except Exception as e:
        return None, None, str(e)


def save_with_original_to_drive(data, source, content, file_path, file_name):
    """Save the original binary file to Drive plus an analysis doc alongside it.
    Large files go through the Drive API directly (no base64); small files use
    the existing Apps Script base64 path."""
    if not APPS_SCRIPT_URL:
        return None, None, None, None, "Apps Script URL not configured"

    folder = data.get("folder", "Knowledge/Market")
    ext = os.path.splitext(file_name)[1].lower()
    mime = MIME_BY_EXT.get(ext, "application/octet-stream")
    analysis_content = (
        f"SUMMARY:\n{data.get('summary','')}\n\nKEY INSIGHTS:\n" +
        "\n".join([f"{i+1}. {t}" for i, t in enumerate(data.get('key_insights', []))]) +
        f"\n\nEXTRACTED CONTENT:\n{content[:6000]}"
    )
    try:
        size = os.path.getsize(file_path)
    except Exception:
        size = 0

    # ---- HYBRID: large file → Drive API direct upload, then register via Apps Script ----
    if size > DRIVE_API_THRESHOLD and GOOGLE_SA_JSON:
        if size > MAX_DRIVE_API_BYTES:
            return None, None, None, None, f"file too large ({size // (1024*1024)} MB > 100 MB)"
        folder_id, ferr = resolve_folder_id(folder)
        if not folder_id:
            return None, None, None, None, f"could not resolve folder for Drive API: {ferr}"
        up_id, up_url, uerr = drive_api_upload(file_path, file_name, folder_id, mime)
        if uerr:
            return None, None, None, None, f"Drive API upload failed: {uerr}"
        # register the uploaded file (analysis doc + indexes) via Apps Script
        try:
            payload = {
                "token": SCRIPT_TOKEN, "action": "register_uploaded",
                "title": data.get("title", "Untitled"), "folder": folder,
                "content": analysis_content, "source": source,
                "importance": data.get("importance", 5), "credibility": data.get("credibility", 5),
                "date": data.get("date", datetime.now().strftime("%Y-%m-%d")),
                "version": data.get("version", "V1.0"), "hashtags": data.get("hashtags", []),
                "summary": data.get("summary", ""), "key_insights": data.get("key_insights", []),
                "original_id": up_id, "original_url": up_url, "file_name": file_name,
            }
            resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=120)
            logger.info(f"register_uploaded response [{resp.status_code}]: {resp.text[:400]}")
            result = resp.json()
            if result.get("ok"):
                return (result.get("file_url"), result.get("file_id"),
                        result.get("original_url") or up_url,
                        result.get("original_id") or up_id, None)
            return None, None, None, None, result.get("error", "register failed")
        except Exception as e:
            return None, None, None, None, f"register after upload failed: {e}"

    # ---- SMALL file → existing base64 path via Apps Script ----
    try:
        import base64
        if size > MAX_ORIGINAL_BYTES:
            return None, None, None, None, f"file too large to store original ({size // (1024*1024)} MB > 30 MB)"
        with open(file_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("ascii")
        payload = {
            "token": SCRIPT_TOKEN,
            "action": "create_with_file",
            "title": data.get("title", "Untitled"),
            "folder": folder,
            "content": analysis_content,
            "source": source,
            "importance": data.get("importance", 5),
            "credibility": data.get("credibility", 5),
            "date": data.get("date", datetime.now().strftime("%Y-%m-%d")),
            "version": data.get("version", "V1.0"),
            "hashtags": data.get("hashtags", []),
            "summary": data.get("summary", ""),
            "key_insights": data.get("key_insights", []),
            "file_b64": b64,
            "file_name": file_name,
            "mime_type": mime,
        }
        resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=300)
        logger.info(f"create_with_file response [{resp.status_code}]: {resp.text[:400]}")
        result = resp.json()
        if result.get("ok"):
            return (result.get("file_url"), result.get("file_id"),
                    result.get("original_url"), result.get("original_id"), None)
        return None, None, None, None, result.get("error", "Unknown error")
    except Exception as e:
        return None, None, None, None, str(e)


def fetch_index_list():
    """Fetch the list of all documents from INDEX via Apps Script."""
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "list_index"
        }, timeout=30)
        try:
            result = resp.json()
        except Exception:
            # not JSON — surface the raw response so we can see what Apps Script returned
            snippet = resp.text[:300].replace("\n", " ")
            logger.info(f"list_index non-JSON [{resp.status_code}]: {snippet}")
            return None, f"[{resp.status_code}] non-JSON response: {snippet}"
        if result.get("ok"):
            return result.get("docs", []), None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


def fetch_docs_content(file_ids):
    """Fetch the full text of documents by a list of IDs."""
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "read_docs", "file_ids": file_ids
        }, timeout=45)
        result = resp.json()
        if result.get("ok"):
            return result.get("docs", []), None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


def pick_relevant_docs(question, index_docs):
    """Step 1: Groq selects relevant documents from the INDEX list."""
    catalog = "\n".join([
        f"{i+1}. [{d.get('folder','')}] {d.get('title','')} (id: {d.get('file_id','')}) {d.get('scores','')}"
        for i, d in enumerate(index_docs) if d.get('file_id')
    ])
    prompt = f"""You are a search assistant for the TrustMe knowledge base. The user question may be in any language — understand it regardless. User question:
"{question}"

Here is the document catalog (number, folder, title, id):
{catalog}

Select up to 3 MOST relevant documents to answer the question.
Return ONLY JSON, no explanation:
{{"file_ids": ["id1", "id2"], "reason": "briefly why"}}
If nothing fits — return {{"file_ids": [], "reason": "..."}}."""

    response = groq_client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.2,
        response_format={"type": "json_object"},
    )
    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```json\s*|^```\s*|\s*```$", "", raw)
    return json.loads(raw)


def answer_from_docs(question, docs_content):
    """Step 2: Groq answers the question using the full text of selected documents."""
    context_text = ""
    for d in docs_content:
        if d.get("content"):
            context_text += f"\n\n=== DOCUMENT: {d.get('title','')} ===\n{d['content']}"

    prompt = f"""You are the TrustMe knowledge base assistant. Answer the user's question RELYING ONLY on the provided documents.
If the documents don't contain the answer — say so honestly, don't make things up.

QUESTION: {question}

DOCUMENTS:{context_text}

Answer in English (even if the documents or question are in another language), to the point. At the end, note which documents you relied on."""

    response = groq_client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.4,
    )
    return response.choices[0].message.content.strip()


async def ask_knowledge_base(update, context, override_text=None):
    """Handle a 'Help, ...' question — search and answer from the knowledge base."""
    if not is_allowed(update):
        await deny(update)
        return
    text = (override_text if override_text is not None else (update.message.text or "")).strip()
    # remove the 'Help' trigger and leading punctuation
    question = re.sub(r"^help[,!\s]*", "", text, flags=re.IGNORECASE).strip()
    if not question:
        await update.message.reply_text(
            "Type a question after «Help», for example:\n"
            "«Help, what materials do we have on KYC?»"
        )
        return

    await update.message.reply_text("🔎 Searching the knowledge base...")

    # Step 1 — get the catalog
    index_docs, error = fetch_index_list()
    if error:
        await update.message.reply_text(f"⚠️ Could not read INDEX: {error}")
        return
    if not index_docs:
        await update.message.reply_text(
            "The index returned no readable rows.\n"
            "If documents exist, their INDEX_Brief rows may be missing a document link, "
            "or the bot points at a different INDEX_Brief than where they were saved."
        )
        return

    # Step 2 — pick relevant
    try:
        pick = pick_relevant_docs(question, index_docs)
    except Exception as e:
        await update.message.reply_text(f"⚠️ Search error: {str(e)[:200]}")
        return

    file_ids = pick.get("file_ids", [])
    if not file_ids:
        await update.message.reply_text(
            "Nothing relevant found in the base for this query.\n"
            f"_{pick.get('reason','')}_",
            parse_mode="Markdown"
        )
        return

    await update.message.reply_text(f"📖 Reading {len(file_ids)} document(s)...")

    # Step 3 — read full text
    docs_content, error = fetch_docs_content(file_ids)
    if error:
        await update.message.reply_text(f"⚠️ Could not read documents: {error}")
        return

    # Step 4 — answer
    try:
        answer = answer_from_docs(question, docs_content)
    except Exception as e:
        await update.message.reply_text(f"⚠️ Answer generation error: {str(e)[:200]}")
        return

    # source references
    sources = "\n".join([
        f"• {d.get('title','')}" for d in docs_content if d.get("content")
    ])
    full = answer
    if len(full) > 3500:
        full = full[:3500] + "..."
    await update.message.reply_text(full, disable_web_page_preview=True)


# ── Voice & Image processing ───────────────────────────────

VISION_MODEL = "meta-llama/llama-4-scout-17b-16e-instruct"  # Groq vision (preview)


def transcribe_audio(file_path):
    """Transcribe an audio file via Groq Whisper (ru/en). Returns (text, error)."""
    try:
        with open(file_path, "rb") as f:
            resp = requests.post(
                "https://api.groq.com/openai/v1/audio/transcriptions",
                headers={"Authorization": f"Bearer {GROQ_KEY}"},
                files={"file": (os.path.basename(file_path), f)},
                data={"model": "whisper-large-v3", "response_format": "text"},
                timeout=90,
            )
        if resp.status_code == 200:
            text = resp.text.strip()
            return (text, None) if text else (None, "empty transcript")
        return None, f"Whisper error {resp.status_code}: {resp.text[:200]}"
    except Exception as e:
        return None, str(e)


def describe_image(file_path):
    """Read/describe an image via Groq vision. Auto-detects text vs photo.
    Returns (text, error)."""
    try:
        import base64
        with open(file_path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("ascii")
        ext = os.path.splitext(file_path)[1].lower().lstrip(".") or "jpeg"
        if ext == "jpg":
            ext = "jpeg"
        prompt = (
            "Look at this image. If it contains text (a document, table, screenshot, "
            "receipt, etc.), extract ALL the text accurately, preserving structure. "
            "If it is a photo, chart or diagram with little text, describe what it shows "
            "in detail. Respond in English. If there is a question or task visible in the "
            "image, note it clearly at the end as 'DETECTED REQUEST: ...'."
        )
        resp = groq_client.chat.completions.create(
            model=VISION_MODEL,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": prompt},
                    {"type": "image_url",
                     "image_url": {"url": f"data:image/{ext};base64,{b64}"}},
                ],
            }],
            temperature=0.3,
        )
        return resp.choices[0].message.content.strip(), None
    except Exception as e:
        return None, str(e)


def detect_intent(text):
    """Decide whether transcribed/extracted text is a QUESTION, a RESEARCH
    request, or MATERIAL to archive. Returns one of: 'question','research','material'."""
    prompt = f"""Classify the user's intent from this message (it may be in Russian or English):

"{text[:1500]}"

Return ONLY JSON:
{{"intent": "question" | "research" | "material"}}

- "question": they are ASKING something that should be answered from our internal knowledge base
- "research": they explicitly want to gather NEW information from the web on a topic
- "material": it is information/content to be SAVED into the knowledge base (a note, fact, document, update)

If unsure, choose "material"."""
    try:
        resp = groq_client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
            response_format={"type": "json_object"},
        )
        raw = re.sub(r"^```json\s*|^```\s*|\s*```$", "", resp.choices[0].message.content.strip())
        return json.loads(raw).get("intent", "material")
    except Exception:
        return "material"


# ── Research feature ───────────────────────────────────────

def tavily_search(query, max_results=4):
    """Search the web via Tavily — returns list of {title, url, content}."""
    if not TAVILY_KEY:
        return None, "Tavily API key not configured"
    try:
        resp = requests.post("https://api.tavily.com/search", json={
            "api_key": TAVILY_KEY,
            "query": query,
            "search_depth": "advanced",
            "max_results": max_results,
            "include_answer": False,
        }, timeout=40)
        resp.raise_for_status()
        data = resp.json()
        results = [{
            "title": r.get("title", ""),
            "url": r.get("url", ""),
            "content": r.get("content", "")
        } for r in data.get("results", [])]
        return results, None
    except Exception as e:
        return None, str(e)


def plan_research_queries(topic, user_material):
    """Groq turns the topic into 2-3 focused web search queries."""
    prompt = f"""You are a research assistant for TrustMe (Islamic digital bank).
The user wants to research this topic: "{topic}"
{f'They also provided this material as context: {user_material[:2000]}' if user_material else ''}

Generate 2-3 focused web search queries (in English) that will find the most useful, factual information.
Return ONLY JSON: {{"queries": ["query 1", "query 2", "query 3"]}}"""
    response = groq_client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
        response_format={"type": "json_object"},
    )
    raw = re.sub(r"^```json\s*|^```\s*|\s*```$", "", response.choices[0].message.content.strip())
    return json.loads(raw).get("queries", [topic])


def synthesize_research(topic, user_material, web_results):
    """Groq writes the FOUND ONLINE section from web results, with inline source refs."""
    sources_block = ""
    for i, r in enumerate(web_results, 1):
        sources_block += f"\n[Source {i}] {r['title']} ({r['url']})\n{r['content'][:2000]}\n"

    prompt = f"""You are a research analyst for TrustMe (Islamic digital bank).
Topic: "{topic}"

You have these web sources:
{sources_block}

Write a clear, factual research summary in ENGLISH based ONLY on these sources.
- Organize by sub-topic with short headings
- After each claim, cite the source number like [Source 1]
- Be concise and factual, no fluff
- If sources conflict, note it
Do NOT invent facts not in the sources."""
    response = groq_client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.4,
    )
    return response.choices[0].message.content.strip()


def pick_research_folder(topic):
    """Groq picks the best folder for the research document."""
    prompt = f"""Topic of a research document: "{topic}"
Available folders:
{folders_for_prompt()}

Pick the single best-fit folder. Return ONLY JSON: {{"folder": "Knowledge/Market"}}"""
    try:
        response = groq_client.chat.completions.create(
            model=MODEL,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2,
            response_format={"type": "json_object"},
        )
        raw = re.sub(r"^```json\s*|^```\s*|\s*```$", "", response.choices[0].message.content.strip())
        folder = json.loads(raw).get("folder", "Knowledge/Market")
        return folder if folder in get_live_folders() else "Knowledge/Market"
    except Exception:
        return "Knowledge/Market"


def build_research_document(topic, user_material, web_results, synthesis):
    """Assemble the final research document body with clear sections."""
    today = datetime.now().strftime("%Y-%m-%d")
    parts = []
    parts.append("RESEARCH DOCUMENT")
    parts.append(f"Topic: {topic}")
    parts.append(f"Date: {today}")
    parts.append("Generated by: TrustMe Knowledge Agent (web research)")
    parts.append("=" * 50)
    parts.append("")
    parts.append("SECTION 1 — FROM USER")
    parts.append("(Material and request provided by the team member)")
    parts.append("")
    parts.append(f"Request: {topic}")
    if user_material:
        parts.append("")
        parts.append("Provided material:")
        parts.append(user_material[:5000])
    else:
        parts.append("(No material attached — topic-only research request)")
    parts.append("")
    parts.append("=" * 50)
    parts.append("")
    parts.append("SECTION 2 — FOUND ONLINE")
    parts.append("(Researched by the agent from open web sources)")
    parts.append("")
    parts.append(synthesis)
    parts.append("")
    parts.append("=" * 50)
    parts.append("")
    parts.append("SECTION 3 — SOURCES")
    parts.append("")
    for i, r in enumerate(web_results, 1):
        parts.append(f"[Source {i}] {r['title']}")
        parts.append(f"    {r['url']}")
    return "\n".join(parts)


def save_research_to_drive(data, source, content):
    """Save a research document via Apps Script create_doc."""
    if not APPS_SCRIPT_URL:
        return None, None, "Apps Script URL not configured"
    payload = {
        "token": SCRIPT_TOKEN, "action": "create_doc",
        "title": data["title"], "folder": data["folder"], "content": content,
        "source": source, "importance": data["importance"], "credibility": data["credibility"],
        "date": data["date"], "version": data["version"], "hashtags": data["hashtags"],
        "summary": data["summary"], "key_insights": data["key_insights"],
    }
    try:
        resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=40)
        logger.info(f"Research save response [{resp.status_code}]: {resp.text[:300]}")
        result = resp.json()
        if result.get("ok"):
            return result.get("file_url"), result.get("file_id"), None
        return None, None, result.get("error", "Unknown error")
    except Exception as e:
        return None, None, str(e)


async def research(update, context, override_text=None):
    """Handle 'Research: <topic>' — web research compiled into a Drive document."""
    if not is_allowed(update):
        await deny(update)
        return
    text = (override_text if override_text is not None else (update.message.text or "")).strip()
    topic = re.sub(r"^research[:,!\s]*", "", text, flags=re.IGNORECASE).strip()
    user_material = context.user_data.get("pending_text", "") or ""

    if not topic:
        await update.message.reply_text(
            "Type a topic after «Research:», for example:\n"
            "«Research: halal ETF competitors and their fee structures»"
        )
        return

    if not TAVILY_KEY:
        await update.message.reply_text(
            "⚠️ Web research is not configured yet (missing Tavily key). "
            "Add TAVILY_API_KEY in Railway and redeploy."
        )
        return

    await update.message.reply_text(f"🔬 Researching: {topic}\nPlanning search queries...")

    try:
        queries = plan_research_queries(topic, user_material)
    except Exception as e:
        await update.message.reply_text(f"⚠️ Could not plan research: {str(e)[:200]}")
        return

    all_results = []
    seen_urls = set()
    for q in queries[:3]:
        await update.message.reply_text(f"🌐 Searching: {q}")
        results, err = tavily_search(q, max_results=4)
        if results:
            for r in results:
                if r["url"] not in seen_urls and r.get("content"):
                    seen_urls.add(r["url"])
                    all_results.append(r)

    if not all_results:
        await update.message.reply_text("No web results found for this topic. Try rephrasing.")
        return

    all_results = all_results[:8]
    await update.message.reply_text(f"📚 Found {len(all_results)} sources. Synthesizing...")

    try:
        synthesis = synthesize_research(topic, user_material, all_results)
    except Exception as e:
        await update.message.reply_text(f"⚠️ Synthesis error: {str(e)[:200]}")
        return

    folder = pick_research_folder(topic)
    body = build_research_document(topic, user_material, all_results, synthesis)
    title = f"Research — {topic[:60]}"

    await update.message.reply_text(f"💾 Saving research to {folder}...")
    payload_data = {
        "title": title, "folder": folder,
        "summary": f"Web research on: {topic}", "key_insights": [],
        "importance": 6, "credibility": 6,
        "date": datetime.now().strftime("%Y-%m-%d"), "version": "V1.0",
        "hashtags": ["#research", "#type_research", "#lang_en"],
    }
    drive_url, file_id, error = save_research_to_drive(payload_data, "Web research", body)

    if drive_url:
        context.bot_data.setdefault("docs", []).append({
            "title": title, "folder": folder, "importance": 6, "credibility": 6,
            "date": datetime.now().strftime("%Y-%m-%d"), "source": "Web research",
            "drive_url": drive_url, "file_id": file_id or "",
        })
        chat_preview = synthesis if len(synthesis) <= 3000 else synthesis[:3000] + "..."
        await update.message.reply_text(
            f"✅ Research saved\n📁 {folder}\n📄 {title}\n🔗 {drive_url}",
            disable_web_page_preview=True
        )
        await update.message.reply_text(chat_preview, disable_web_page_preview=True)
        await notify_admins(context, update,
            f"📢 {_display_name(update)} ran web research\n"
            f"🔬 {topic}\n📁 {folder}\n🔗 {drive_url}")
        context.user_data.clear()
    else:
        await update.message.reply_text(f"⚠️ Could not save research: {error}")


def esc(text):
    """Escape Telegram Markdown (legacy) special chars in dynamic text."""
    if text is None:
        return ""
    for ch in ("_", "*", "`", "["):
        text = str(text).replace(ch, "\\" + ch)
    return text


def format_preview(data, source):
    tags     = "  ".join(data.get("hashtags", []))
    insights = "\n".join([f"  {i+1}. {esc(t)}" for i, t in enumerate(data.get("key_insights", []))])
    imp      = data.get("importance",  0)
    cred     = data.get("credibility", 0)
    return (
        f"📄 *{esc(data.get('title','Untitled'))}*\n\n"
        f"📁 *Folder:* `{data.get('folder','—')}`\n"
        f"📅 {data.get('date','—')}  |  🏷 {data.get('version','V1.0')}\n\n"
        f"─────────────────────\n"
        f"📝 *Summary*\n{esc(data.get('summary','—'))}\n\n"
        f"💡 *Key Insights*\n{insights}\n\n"
        f"🏷 *Hashtags*\n{esc(tags)}\n\n"
        f"─────────────────────\n"
        f"⭐️ *Importance:* {imp}/10  {'🟠'*imp}{'⚪️'*(10-imp)}\n"
        f"_{esc(data.get('importance_reason',''))}_\n\n"
        f"✅ *Credibility:* {cred}/10  {'🔵'*cred}{'⚪️'*(10-cred)}\n"
        f"_{esc(data.get('credibility_reason',''))}_\n\n"
        f"🔗 *Source:* {esc(source)}"
    )


# ── Command handlers ──────────────────────────────────────

async def start(update, context):
    if not is_allowed(update):
        await deny(update)
        return
    register_user(update)
    await update.message.reply_text(
        "👋 *NeoBank Knowledge Agent*\n"
        "_TrustMe Knowledge Base · Groq + Google Drive_\n\n"
        "I help collect and store the team's knowledge in Google Drive.\n\n"
        "*What I can do:*\n"
        "📎 Accept a link, file, or text\n"
        "🧠 Analyze and sort it into folders\n"
        "🗂 Update indexes and changelog\n"
        "↩️ Roll back a mistaken save\n\n"
        "*How to start:* just send me a link, file, or text.\n\n"
        "💡 To ask the knowledge base — start your message with the word «Help»:\n"
        "_«Help, what materials do we have on KYC?»_\n\n"
        "🔬 To run web research — start with «Research:»:\n"
        "_«Research: halal ETF competitors and fees»_\n\n"
        "Full guide — /help\n"
        "Command list — the «☰» or «/» button next to the input field.",
        parse_mode="Markdown"
    )


async def help_cmd(update, context):
    if not is_allowed(update):
        return
    await update.message.reply_text(
        "📖 *NeoBank Knowledge Agent — Guide*\n\n"
        "*1. How to add material*\n"
        "Send the bot:\n"
        "🔗 a link — article, page, post\n"
        "📎 a file — PDF, DOCX, Excel (XLSX/XLS), CSV, ODS, TXT\n"
        "💬 text — just type or paste\n\n"
        "*2. What happens*\n"
        "The bot will ask what to focus on. Then two paths:\n"
        "• Type a focus or /skip → the bot analyzes via AI, "
        "suggests a folder, tags and scores, shows a preview\n"
        "• «📥 Save without analysis» button → pick a folder, "
        "the file is stored as-is, unprocessed\n\n"
        "*3. Confirmation*\n"
        "In the preview you can: save, change folder, adjust scores, or discard.\n"
        "After saving, the document goes to the right Drive folder, "
        "and the indexes and changelog update automatically.\n\n"
        "*4. Rollback*\n"
        "/undo — roll back the last save\n"
        "/rollback <ID> — roll back a document by ID (from a link or /list)\n"
        "Rolled-back documents go to _Archive — they are not deleted and can be restored.\n\n"
        "*5. Questions to the knowledge base*\n"
        "Start your message with the word «Help» — the bot finds the relevant documents and answers from their content:\n"
        "_«Help, what potential KYC partners do we have?»_\n"
        "_«Help, find materials on MoonPay»_\n\n"
        "*6. Web research*\n"
        "Start your message with «Research:» — the bot searches the open web, "
        "compiles findings with source links, and saves a document split into "
        "FROM USER / FOUND ONLINE / SOURCES:\n"
        "_«Research: Sharia-compliant stablecoin regulation in UAE»_\n\n"
        "*7. Commands*\n"
        "/start — welcome\n"
        "/help — this guide\n"
        "/status — knowledge base stats\n"
        "/list — recent documents\n"
        "/undo — roll back the last save\n"
        "/rollback — roll back by ID\n\n"
        "*Storage rules:*\n"
        "• Old content is never deleted — it moves to _Archive\n"
        "• Every action is logged in CHANGELOG\n"
        "• File name: Title_Version_Date",
        parse_mode="Markdown"
    )

async def status(update, context):
    if not is_allowed(update):
        return
    await update.message.reply_text("📊 Reading the knowledge base...")

    index_docs, error = fetch_index_list()
    if error:
        await update.message.reply_text(f"⚠️ Could not read INDEX: {error}")
        return
    # only count real registry rows (those carrying a folder)
    docs = [d for d in (index_docs or []) if d.get("folder")]
    total = len(docs)
    if not total:
        await update.message.reply_text("📊 No documents in the knowledge base yet. Send your first link!")
        return

    # parse Imp/Cred from the 'scores' field (e.g. "Imp: 6  Cred: 8")
    imp_vals, cred_vals = [], []
    folders = {}
    for d in docs:
        folders[d.get("folder", "Unknown")] = folders.get(d.get("folder", "Unknown"), 0) + 1
        s = d.get("scores", "")
        mi = re.search(r"Imp:\s*(\d+)", s)
        mc = re.search(r"Cred:\s*(\d+)", s)
        if mi: imp_vals.append(int(mi.group(1)))
        if mc: cred_vals.append(int(mc.group(1)))

    avg_imp  = round(sum(imp_vals) / len(imp_vals), 1) if imp_vals else "—"
    avg_cred = round(sum(cred_vals) / len(cred_vals), 1) if cred_vals else "—"
    fl = "\n".join([f"  • {f} — {c}" for f, c in sorted(folders.items())])

    await update.message.reply_text(
        f"📊 Knowledge Base\n\n"
        f"📄 {total} documents\n"
        f"🗂 {len(folders)} folders\n"
        f"⭐️ Avg importance: {avg_imp}/10\n"
        f"✅ Avg credibility: {avg_cred}/10\n\n"
        f"📁 By folder:\n{fl}",
        disable_web_page_preview=True
    )

async def newfolder_cmd(update, context):
    """/newfolder <Parent/Path> <Name> — admins only."""
    if not is_allowed(update):
        await deny(update)
        return
    args = context.args
    if len(args) < 2:
        await update.message.reply_text(
            "Usage: /newfolder <Parent/Path> <Name>\n\n"
            "Examples:\n"
            "• /newfolder Company/Banking_Ops Cards\n"
            "• /newfolder Knowledge Legal\n\n"
            "Name must be English letters, digits or underscore (no spaces)."
        )
        return
    parent_path, name = args[0], args[1]
    created_by = (update.effective_user.username or "unknown").lower()
    await update.message.reply_text(f"📁 Creating folder «{name}» in {parent_path}...")
    path, error = create_folder_in_drive(parent_path, name, created_by)
    if path:
        invalidate_folder_cache()
        await update.message.reply_text(f"✅ Folder created: {path}\nIt's now available when saving.")
        await notify_admins(context, update,
            f"📢 {_display_name(update)} created a folder\n📁 {path}")
    else:
        await update.message.reply_text(f"⚠️ Could not create folder: {error}")


async def testnotify_cmd(update, context):
    """/testnotify — send a test notification to the other admins and report the result."""
    if not is_allowed(update):
        await deny(update)
        return
    register_user(update)
    diag = await notify_admins(context, update,
        f"🔔 Test notification from {_display_name(update)}. If you see this, notifications work.")
    await update.message.reply_text(
        f"Test sent.\nResult: {diag}\n\n"
        "If it says 'notified 0' with no errors, only you are registered right now, "
        "or the others are the ones who should receive it. Ask another admin to run /testnotify too."
    )


async def whoisregistered_cmd(update, context):
    """/whoisregistered — show who is on the notification list (diagnostic)."""
    if not is_allowed(update):
        await deny(update)
        return
    # ensure the caller is registered right now
    register_user(update)
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "list_users"
        }, timeout=10)
        users = resp.json().get("users", {})
    except Exception as e:
        await update.message.reply_text(f"⚠️ Could not read the list: {e}")
        return
    if not users:
        await update.message.reply_text(
            "📭 No one is registered for notifications yet.\n\n"
            "Each admin must send the bot a message (e.g. /start) once so the bot "
            "learns their chat. Ask the team to do that, then check again."
        )
        return
    lines = "\n".join([f"• {u}" for u in users.keys()])
    await update.message.reply_text(
        f"🔔 Registered for notifications ({len(users)}):\n{lines}"
    )


async def drivetest_cmd(update, context):
    """/drivetest — verify direct Google Drive API access (migration check)."""
    if not is_allowed(update):
        await deny(update)
        return
    await update.message.reply_text("🔌 Testing direct Google Drive API connection...")
    # show which Service Account identity we're using
    try:
        import json as _json
        info = _json.loads(GOOGLE_SA_JSON) if GOOGLE_SA_JSON else {}
        sa_email = info.get("client_email", "(no client_email in key)")
        await update.message.reply_text(
            f"🔑 Using Service Account:\n`{sa_email}`\n\n"
            "This exact email must be a MEMBER of the Shared Drive (via Manage Members of the "
            "drive itself, not folder sharing).",
            parse_mode="Markdown"
        )
    except Exception as e:
        await update.message.reply_text(f"⚠️ Could not read key: {e}")
    result = drive_api_selftest()
    await update.message.reply_text(result, disable_web_page_preview=True)


async def setscore_cmd(update, context):
    """/setscore <doc-id or "exact title"> <importance> <credibility>
    Updates a document's scores in the registry so /toc keeps them."""
    if not is_allowed(update):
        await deny(update)
        return
    text = (update.message.text or "").strip()
    # strip the command word
    rest = text.split(None, 1)[1] if len(text.split(None, 1)) > 1 else ""
    # scores are the last two integers; key is everything before them
    m = re.match(r'^(.*?)(?:\s+)(\d+)\s+(\d+)\s*$', rest)
    if not m:
        await update.message.reply_text(
            "Usage:\n`/setscore <doc-id or exact title> <importance> <credibility>`\n\n"
            "Examples:\n"
            "`/setscore 1AbC...xyz 8 9`  (by document ID from its link)\n"
            "`/setscore Fasset_Analysis_V1.0_2026-08-17 8 9`  (by exact title)",
            parse_mode="Markdown"
        )
        return
    key = m.group(1).strip().strip('"').strip("'")
    imp, cred = int(m.group(2)), int(m.group(3))
    if not (1 <= imp <= 10 and 1 <= cred <= 10):
        await update.message.reply_text("⚠️ Both scores must be 1–10.")
        return
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "set_score",
            "key": key, "importance": imp, "credibility": cred,
        }, timeout=30)
        result = resp.json()
    except Exception as e:
        await update.message.reply_text(f"⚠️ Error: {str(e)[:200]}")
        return
    if result.get("ok"):
        await update.message.reply_text(
            f"✅ Scores updated\n📄 {result.get('title', key)}\n"
            f"⭐️ Importance {imp}/10   ✅ Credibility {cred}/10\n\n"
            "Run /toc to refresh the table of contents."
        )
    else:
        await update.message.reply_text(
            f"⚠️ Could not update: {result.get('error','unknown')}\n\n"
            "Tip: use the document's exact title, or its ID from the /d/<ID> part of its link."
        )


async def toc_cmd(update, context):
    """/toc — rebuild the table of contents in both index documents."""
    if not is_allowed(update):
        await deny(update)
        return
    await update.message.reply_text("📑 Rebuilding the table of contents...")
    if not APPS_SCRIPT_URL:
        await update.message.reply_text("⚠️ Apps Script not configured.")
        return
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "rebuild_toc"
        }, timeout=180)
        try:
            result = resp.json()
        except Exception:
            snippet = resp.text[:300].replace("\n", " ")
            await update.message.reply_text(
                f"⚠️ Apps Script returned non-JSON [{resp.status_code}].\n"
                f"This means the new script version isn't deployed.\n\n"
                f"Raw start: {snippet}"
            )
            return
        if result.get("ok"):
            await update.message.reply_text(
                f"✅ Table of contents rebuilt.\n{result.get('result','')}\n\n"
                "Check INDEX_Brief (short) and INDEX_Detailed (full)."
            )
        else:
            await update.message.reply_text(f"⚠️ Could not rebuild: {result.get('error','unknown')}")
    except Exception as e:
        await update.message.reply_text(f"⚠️ Error: {str(e)[:200]}")


async def delfolder_cmd(update, context):
    """/delfolder <Full/Path> — archives a folder (superadmin any, admin only own)."""
    if not is_allowed(update):
        await deny(update)
        return
    args = context.args
    if not args:
        await update.message.reply_text(
            "Usage: /delfolder <Full/Path>\n\n"
            "Example: /delfolder Company/Banking_Ops/Cards\n\n"
            "The folder and its contents move to _Archive (nothing is deleted).\n"
            "You can archive folders you created; the superadmin can archive any."
        )
        return
    path = " ".join(args).strip()
    requested_by = (update.effective_user.username or "unknown").lower()
    await update.message.reply_text(f"🗑 Archiving folder {path}...")
    result, error = delete_folder_in_drive(path, requested_by, is_superadmin(update))
    if result:
        invalidate_folder_cache()
        await update.message.reply_text(
            f"✅ Folder moved to _Archive: {path}\n"
            f"Nothing is deleted — it can be restored from _Archive."
        )
        await notify_admins(context, update,
            f"📢 {_display_name(update)} archived a folder\n📁 {path} → _Archive")
    else:
        await update.message.reply_text(f"⚠️ Could not archive folder: {error}")


async def list_docs(update, context):
    if not is_allowed(update):
        return
    docs = context.bot_data.get("docs", [])
    if not docs:
        await update.message.reply_text("No documents yet.")
        return
    lines = []
    for i, d in enumerate(docs[-10:][::-1], 1):
        drive_link = f" · [Drive]({d.get('drive_url','')})" if d.get('drive_url') else ""
        lines.append(
            f"{i}. *{d.get('title','Untitled')}*\n"
            f"   `{d.get('folder','—')}` | ⭐️{d.get('importance','—')} ✅{d.get('credibility','—')} | {d.get('date','—')}{drive_link}"
        )
    await update.message.reply_text(
        "📋 *Recent Documents*\n\n" + "\n\n".join(lines),
        parse_mode="Markdown", disable_web_page_preview=True
    )


def save_raw_to_drive(folder, title, text_content, source):
    """Save a document to Drive without AI analysis."""
    if not APPS_SCRIPT_URL:
        return None, "Apps Script URL not configured"
    payload = {
        "token": SCRIPT_TOKEN,
        "action": "create_doc",
        "title": title,
        "folder": folder,
        "content": text_content or "(file saved without processing)",
        "source": source,
        "importance": 0,
        "credibility": 0,
        "date": datetime.now().strftime("%Y-%m-%d"),
        "version": "V1.0",
        "hashtags": ["#raw_unprocessed"],
        "summary": "Saved without analysis",
        "key_insights": [],
    }
    try:
        resp = requests.post(APPS_SCRIPT_URL, json=payload, timeout=30)
        logger.info(f"Raw save response [{resp.status_code}]: {resp.text[:300]}")
        result = resp.json()
        if result.get("ok"):
            return result, None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


def rollback_in_drive(file_id):
    """Ask Apps Script to move a document to _Archive."""
    if not APPS_SCRIPT_URL:
        return None, "Apps Script URL not configured"
    try:
        resp = requests.post(APPS_SCRIPT_URL, json={
            "token": SCRIPT_TOKEN, "action": "rollback", "file_id": file_id
        }, timeout=30)
        logger.info(f"Rollback response [{resp.status_code}]: {resp.text[:300]}")
        result = resp.json()
        if result.get("ok"):
            return result, None
        return None, result.get("error", "Unknown error")
    except Exception as e:
        return None, str(e)


async def undo(update, context):
    """Roll back the last created document — move it to _Archive."""
    if not is_allowed(update):
        return
    docs = context.bot_data.get("docs", [])
    # find the last document with a file_id that hasn't been rolled back yet
    target = None
    for d in reversed(docs):
        if d.get("file_id") and not d.get("rolled_back"):
            target = d
            break
    if not target:
        await update.message.reply_text("Nothing to roll back — no recent saves with an ID.")
        return

    await update.message.reply_text(f"↩️ Rolling back: {target.get('title','Untitled')}...")
    result, error = rollback_in_drive(target["file_id"])
    if result:
        target["rolled_back"] = True
        extra = ""
        # if this save had an original file (Excel/PDF etc.), archive it too
        orig_id = target.get("original_id")
        if orig_id:
            r2, e2 = rollback_in_drive(orig_id)
            extra = "\n📎 Original file also moved to _Archive." if r2 else \
                    f"\n⚠️ Analysis archived, but original file failed: {e2}"
        await update.message.reply_text(
            f"✅ Rolled back to _Archive\n"
            f"📄 {result.get('title', target.get('title'))}\n"
            f"📁 {result.get('from','—')} → _Archive"
            f"{extra}\n\n"
            f"Nothing is deleted — items can be restored from _Archive.",
            disable_web_page_preview=True
        )
        await notify_admins(context, update,
            f"📢 {_display_name(update)} rolled back a document\n"
            f"📄 {result.get('title', target.get('title'))} → _Archive")
    else:
        await update.message.reply_text(f"⚠️ Could not roll back: {error}")


async def rollback(update, context):
    """Roll back a document by ID: /rollback <file_id>"""
    if not is_allowed(update):
        return
    args = context.args
    if not args:
        await update.message.reply_text(
            "Usage: /rollback <document ID>\n\n"
            "The ID can be taken from the document link in INDEX or from /list."
        )
        return
    raw = args[0].strip()
    # remove angle brackets, quotes, spaces
    raw = raw.strip("<>\"' ")
    # if a full link was pasted — extract the ID from /d/.../ or ?id=
    m = re.search(r"/d/([a-zA-Z0-9_-]+)", raw) or re.search(r"[?&]id=([a-zA-Z0-9_-]+)", raw)
    file_id = m.group(1) if m else raw
    await update.message.reply_text(f"↩️ Rolling back document {file_id}...")
    result, error = rollback_in_drive(file_id)
    if result:
        extra = ""
        # mark in history and archive the linked original file if we know it
        for d in context.bot_data.get("docs", []):
            if d.get("file_id") == file_id:
                d["rolled_back"] = True
                if d.get("original_id"):
                    r2, e2 = rollback_in_drive(d["original_id"])
                    extra = "\n📎 Original file also moved to _Archive." if r2 else \
                            f"\n⚠️ Analysis archived, but original file failed: {e2}"
                break
        await update.message.reply_text(
            f"✅ Rolled back to _Archive\n"
            f"📄 {result.get('title','—')}\n"
            f"📁 {result.get('from','—')} → _Archive"
            f"{extra}\n\n"
            f"Nothing is deleted — items can be restored from _Archive.",
            disable_web_page_preview=True
        )
        await notify_admins(context, update,
            f"📢 {_display_name(update)} rolled back a document\n"
            f"📄 {result.get('title','—')} → _Archive")
    else:
        await update.message.reply_text(f"⚠️ Could not roll back: {error}")


# ── Receive material ──────────────────────────────────────

async def handle_voice(update, context):
    """Voice message: transcribe → detect intent → answer or save."""
    if not is_allowed(update):
        await deny(update)
        return
    register_user(update)
    msg = update.message
    await msg.reply_text("🎙 Transcribing voice message...")

    tg_file = await context.bot.get_file(msg.voice.file_id)
    fd, tmp_path = tempfile.mkstemp(suffix=".ogg")
    os.close(fd)
    await tg_file.download_to_drive(tmp_path)

    transcript, error = transcribe_audio(tmp_path)
    if error or not transcript:
        try: os.unlink(tmp_path)
        except Exception: pass
        await msg.reply_text(f"⚠️ Couldn't transcribe: {error or 'empty'}")
        return ConversationHandler.END

    await msg.reply_text(f"📝 Transcript:\n{transcript[:1500]}")
    intent = detect_intent(transcript)

    if intent == "question":
        await msg.reply_text("💡 Sounds like a question — searching the knowledge base...")
        try: os.unlink(tmp_path)
        except Exception: pass
        await ask_knowledge_base(update, context, override_text="Help, " + transcript)
        return ConversationHandler.END
    if intent == "research":
        await msg.reply_text("🔬 Sounds like a research request — gathering sources...")
        try: os.unlink(tmp_path)
        except Exception: pass
        await research(update, context, override_text="Research: " + transcript)
        return ConversationHandler.END

    await msg.reply_text("💾 Saving as material. What should I focus on? Send focus or /skip.")
    stored = make_filename("Voice_note")
    context.user_data.update({
        "pending_source": "Voice message", "pending_type": "voice_material",
        "pending_text": transcript, "pending_original_path": tmp_path,
        "pending_original_name": stored + ".ogg",
    })
    return WAITING_FOCUS


async def handle_photo(update, context):
    """Photo/image: read via vision → detect if it's a request → answer or save."""
    if not is_allowed(update):
        await deny(update)
        return
    register_user(update)
    msg = update.message
    await msg.reply_text("🖼 Reading the image...")

    if msg.photo:
        file_id = msg.photo[-1].file_id
        ext = ".jpg"
    else:
        # image sent as a document — check size (photos compressed by Telegram are always small)
        if (getattr(msg.document, "file_size", None) or 0) > TELEGRAM_DOWNLOAD_LIMIT:
            await msg.reply_text(
                "⚠️ This image is over 20 MB — Telegram won't let me download it. "
                "Please compress it and resend."
            )
            return
        file_id = msg.document.file_id
        ext = os.path.splitext(msg.document.file_name or "img.jpg")[1] or ".jpg"

    tg_file = await context.bot.get_file(file_id)
    fd, tmp_path = tempfile.mkstemp(suffix=ext)
    os.close(fd)
    await tg_file.download_to_drive(tmp_path)

    extracted, error = describe_image(tmp_path)
    if error or not extracted:
        try: os.unlink(tmp_path)
        except Exception: pass
        await msg.reply_text(f"⚠️ Couldn't read the image: {error or 'empty'}")
        return ConversationHandler.END

    await msg.reply_text(f"📄 Extracted:\n{extracted[:1500]}", disable_web_page_preview=True)

    if "DETECTED REQUEST:" in extracted:
        req = extracted.split("DETECTED REQUEST:", 1)[1].strip()
        if req:
            await msg.reply_text("💡 The image contains a request — searching the knowledge base...")
            await ask_knowledge_base(update, context, override_text="Help, " + req)

    await msg.reply_text("💾 Saving the image as material. Send focus or /skip.")
    stored = make_filename("Image")
    context.user_data.update({
        "pending_source": "Image", "pending_type": "image_material",
        "pending_text": extracted, "pending_original_path": tmp_path,
        "pending_original_name": stored + ext,
    })
    return WAITING_FOCUS


async def receive_message(update, context):
    if not is_allowed(update):
        await deny(update)
        return
    register_user(update)
    msg = update.message
    if msg.document or msg.audio:
        f = msg.document or msg.audio
        # Telegram Bot API cannot download files larger than 20 MB — check up front
        file_size = getattr(f, "file_size", None) or 0
        if file_size > TELEGRAM_DOWNLOAD_LIMIT:
            size_mb = file_size / (1024 * 1024)
            await msg.reply_text(
                f"⚠️ This file is {size_mb:.0f} MB. Telegram doesn't allow bots to download "
                f"files larger than 20 MB, so I can't process it directly.\n\n"
                f"Please compress it first and send the smaller version:\n"
                f"👉 https://www.ilovepdf.com/ru/compress_pdf\n\n"
                f"(For non-PDF files, you can also upload the file to Google Drive manually "
                f"and send me the link instead.)",
                disable_web_page_preview=True
            )
            return
        context.user_data.update({"pending_source": f"File: {f.file_name}", "pending_type": "file",
                                   "pending_file_id": f.file_id, "pending_file_name": f.file_name,
                                   "pending_file_size": file_size})
    elif msg.voice:
        # voice notes are small, but guard anyway
        if (getattr(msg.voice, "file_size", None) or 0) > TELEGRAM_DOWNLOAD_LIMIT:
            await msg.reply_text("⚠️ This voice message is too large for me to download (over 20 MB).")
            return
        context.user_data.update({"pending_source": "Voice message", "pending_type": "voice",
                                   "pending_file_id": msg.voice.file_id})
    elif msg.text:
        t = msg.text.strip()
        if re.match(r"https?://\S+", t):
            context.user_data.update({"pending_source": t, "pending_type": "url", "pending_url": t})
        else:
            context.user_data.update({"pending_source": "Direct text", "pending_type": "text", "pending_text": t})
    else:
        await msg.reply_text("Please send a link, file, or text.")
        return
    kb = [[InlineKeyboardButton("📥 Save without analysis", callback_data="raw_save")]]
    await msg.reply_text(
        "🎯 *What should I focus on?*\n\n"
        "Tell me what's important for TrustMe, or /skip for default analysis.\n"
        "Or save the file as-is, without processing:",
        parse_mode="Markdown",
        reply_markup=InlineKeyboardMarkup(kb)
    )
    return WAITING_FOCUS


async def receive_focus(update, context):
    if not is_allowed(update):
        return ConversationHandler.END
    # if a raw file is waiting for scores, handle that instead of treating text as focus
    if context.user_data.get("raw_pending"):
        await handle_score_edit(update, context)
        return ConversationHandler.END
    focus = "" if update.message.text.strip().startswith("/skip") else update.message.text.strip()
    await update.message.reply_text("⏳ Analysing...")

    ptype  = context.user_data.get("pending_type")
    source = context.user_data.get("pending_source", "Unknown")
    content = ""

    try:
        if ptype == "url":
            url = context.user_data["pending_url"]
            # Video links can't be read — bot has no transcript access
            if is_video_url(url):
                await update.message.reply_text(
                    "🎬 This looks like a video link. I can't read video or audio content — "
                    "I only process web articles, PDFs, DOCX and text.\n\n"
                    "What you can do:\n"
                    "• Paste the video's transcript or description as text\n"
                    "• Send an article about the same topic\n"
                    "• Use «Research: <topic>» to gather written sources on it"
                )
                return ConversationHandler.END
            await update.message.reply_text("🌐 Fetching URL...")
            content = fetch_url_content(url)
        elif ptype == "text":
            content = context.user_data["pending_text"]
        elif ptype in ("voice_material", "image_material"):
            # transcript / image description already extracted; original file kept
            content = context.user_data.get("pending_text", "")
        elif ptype == "file":
            tg_file = await context.bot.get_file(context.user_data["pending_file_id"])
            orig_name = context.user_data.get("pending_file_name", "file.txt")
            ext = os.path.splitext(orig_name)[1] or ".bin"
            fd, tmp_path = tempfile.mkstemp(suffix=ext)
            os.close(fd)
            await tg_file.download_to_drive(tmp_path)
            content = extract_file_content(tmp_path, orig_name)
            # for structural formats, keep the original file until save; else delete now
            if is_structural_file(orig_name) and content:
                context.user_data["pending_original_path"] = tmp_path
                context.user_data["pending_original_name"] = orig_name
            else:
                os.unlink(tmp_path)

        # fetch/extract failed or returned nothing usable — stop, don't save junk.
        # voice/image/text already have validated content; use a small floor for them.
        floor = 3 if ptype in ("voice_material", "image_material", "text") else 50
        if not content or len(content.strip()) < floor:
            if ptype == "url":
                await update.message.reply_text(
                    "⚠️ I couldn't read that page. It may block automated access, require login, "
                    "or be a video/app page with no article text.\n\n"
                    "Try pasting the text directly, or send a different link."
                )
            else:
                detail = getattr(extract_file_content, "last_error", "") or "unknown"
                await update.message.reply_text(
                    "⚠️ Couldn't extract readable content from that file.\n"
                    f"Reason: {detail}\n\n"
                    "Supported: PDF, DOCX, XLSX, XLS, CSV, ODS, TXT. Or paste the text."
                )
            return ConversationHandler.END

        # guard against unreadable markup/binary that would make the AI hallucinate
        if ptype not in ("voice_material", "image_material") and looks_like_garbage(content):
            await update.message.reply_text(
                "⚠️ I could open the file but the content came out as formatting codes, "
                "not readable text — so I won't guess at a summary.\n\n"
                "This usually means the format isn't fully supported. Try exporting to "
                "PDF, DOCX or plain TXT, or paste the text directly."
            )
            return ConversationHandler.END

        data = process_with_ai(content, source, focus)
        context.user_data["pending_data"]    = data
        context.user_data["pending_content"] = content

        keyboard = [
            [InlineKeyboardButton("✅ Save to Drive", callback_data="confirm_save"),
             InlineKeyboardButton("✏️ Edit scores",   callback_data="edit_scores")],
            [InlineKeyboardButton("🔄 Change folder", callback_data="change_folder"),
             InlineKeyboardButton("❌ Discard",        callback_data="discard")],
        ]
        preview = format_preview(data, source) + "\n\n─────────────────────\n✅ Save to Google Drive?"
        try:
            await update.message.reply_text(
                preview, parse_mode="Markdown",
                reply_markup=InlineKeyboardMarkup(keyboard)
            )
        except Exception:
            # Markdown failed — resend as plain text so the buttons still work
            plain = re.sub(r"[*_`\[\]]", "", preview)
            await update.message.reply_text(
                plain, reply_markup=InlineKeyboardMarkup(keyboard)
            )

    except json.JSONDecodeError:
        await update.message.reply_text("⚠️ Processing error. Please try again.")
    except Exception as e:
        logger.error(e)
        await update.message.reply_text(f"⚠️ Error: {str(e)[:200]}")

    return ConversationHandler.END


# ── Callback handlers ─────────────────────────────────────

async def _finish_raw(query, context, folder, title, result, error):
    """Record and report a raw (no-analysis) save of text/url/non-structural file."""
    if result and result.get("ok"):
        context.bot_data.setdefault("docs", []).append({
            "title": title, "folder": folder, "importance": "—", "credibility": "—",
            "date": datetime.now().strftime("%Y-%m-%d"), "source": query.from_user.username or "—",
            "drive_url": result.get("file_url", ""), "file_id": result.get("file_id", ""),
            "original_id": "",
        })
        await query.message.reply_text(
            f"✅ Saved without analysis\n📁 {folder}\n📄 {title}\n🔗 {result.get('file_url','')}",
            disable_web_page_preview=True
        )
        await notify_admins(context, query,
            f"📢 {_display_name(query)} saved a note (no analysis)\n"
            f"📄 {title}\n📁 {folder}\n🔗 {result.get('file_url','')}")
    else:
        await query.message.reply_text(f"⚠️ Could not save: {error}")
    context.user_data.clear()


async def button_callback(update, context):
    query = update.callback_query
    await query.answer()
    pd  = context.user_data.get("pending_data", {})
    pc  = context.user_data.get("pending_content", "")
    src = context.user_data.get("pending_source", "Unknown")

    if query.data == "confirm_save":
        orig_path = context.user_data.get("pending_original_path")
        orig_name = context.user_data.get("pending_original_name")
        title  = pd.get("title", "Untitled")
        folder = pd.get("folder", "Knowledge/Market")
        original_url = ""
        original_id = ""

        if orig_path and os.path.exists(orig_path):
            # structural file — save original + analysis doc alongside
            await query.message.reply_text("💾 Saving original file and analysis to Google Drive...")
            drive_url, file_id, original_url, original_id, error = save_with_original_to_drive(
                pd, src, pc, orig_path, orig_name
            )
            try:
                os.unlink(orig_path)
            except Exception:
                pass
        else:
            await query.message.reply_text("💾 Saving to Google Drive...")
            drive_url, file_id, error = save_to_drive(pd, src, pc)

        context.bot_data.setdefault("docs", []).append({
            "title":       title,
            "folder":      folder,
            "importance":  pd.get("importance",  5),
            "credibility": pd.get("credibility", 5),
            "date":        pd.get("date", datetime.now().strftime("%Y-%m-%d")),
            "source":      src,
            "drive_url":   drive_url or "",
            "file_id":     file_id or "",
            "original_id": original_id or "",
        })

        if drive_url:
            msg = (
                f"✅ Saved to Google Drive!\n\n"
                f"📁 {folder}\n"
                f"📄 {title}\n"
                f"🔗 Analysis: {drive_url}\n"
            )
            if original_url:
                msg += f"📎 Original file: {original_url}\n"
            msg += "\nINDEX and CHANGELOG updated."
            await query.message.reply_text(msg, disable_web_page_preview=True)
            await notify_admins(context, update,
                f"📢 {_display_name(update)} added a document\n"
                f"📄 {title}\n📁 {folder}\n🔗 {drive_url}")
        else:
            await query.message.reply_text(
                f"⚠️ Drive save failed: {error}\n\n"
                f"Document processed but not saved to Drive."
            )

        await query.edit_message_reply_markup(reply_markup=None)
        context.user_data.clear()

    elif query.data == "discard":
        # clean up any kept original file
        op = context.user_data.get("pending_original_path")
        if op and os.path.exists(op):
            try: os.unlink(op)
            except Exception: pass
        await query.edit_message_text("❌ Discarded.")
        context.user_data.clear()

    elif query.data == "edit_scores":
        await query.message.reply_text(
            f"✏️ Send: `importance 8` or `credibility 6`\n\n"
            f"Current — ⭐️{pd.get('importance','—')}/10  ✅{pd.get('credibility','—')}/10",
            parse_mode="Markdown"
        )

    elif query.data == "change_folder":
        kb = build_folder_menu("folder", is_allowed(update), "newfolder_analysis")
        await query.message.reply_text("📁 *Choose a folder* (tap a department to see its sub-folders):",
                                        parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(kb))

    elif query.data == "raw_save":
        kb = build_folder_menu("rawfolder", is_allowed(update), "newfolder_raw")
        await query.message.reply_text(
            "📁 *Choose a folder to save the file* (tap a department for sub-folders):",
            parse_mode="Markdown", reply_markup=InlineKeyboardMarkup(kb)
        )

    elif query.data.startswith("expand_"):
        # user tapped a top-level folder — show its sub-folders (or pick it directly)
        _, prefix, top = query.data.split("|", 2)  # expand_|folder|Company/Banking_Ops
        subs = [f for f in get_live_folders() if f.startswith(top + "/")]
        rows = [[InlineKeyboardButton(f"📂 {top}  (save here)", callback_data=f"{prefix}_{top}")]]
        for s in subs:
            leaf = s[len(top) + 1:]
            rows.append([InlineKeyboardButton(f"   └ {leaf}", callback_data=f"{prefix}_{s}")])
        rows.append([InlineKeyboardButton("⬅️ Back", callback_data=("change_folder" if prefix == "folder" else "raw_save"))])
        await query.edit_message_text(f"📁 {top} — choose sub-folder or save at department level:",
                                       reply_markup=InlineKeyboardMarkup(rows))

    elif query.data in ("newfolder_analysis", "newfolder_raw"):
        await query.message.reply_text(
            "➕ To create a folder, use the command:\n"
            "/newfolder <Parent/Path> <Name>\n\n"
            "Examples:\n"
            "• /newfolder Company/Banking_Ops Cards\n"
            "• /newfolder Knowledge Legal\n\n"
            "Then send the file again and pick the new folder."
        )

    elif query.data.startswith("rawfolder_"):
        folder = query.data[len("rawfolder_"):]

        ptype = context.user_data.get("pending_type")
        fname = context.user_data.get("pending_file_name", "")

        # build the payload depending on what was sent
        if ptype == "text":
            text_content = context.user_data.get("pending_text", "")
            title = (text_content[:50] + "...") if len(text_content) > 50 else (text_content or "Note")
            result, error = save_raw_to_drive(folder, title, text_content, src)
            await _finish_raw(query, context, folder, title, result, error)
            return

        if ptype == "url":
            url = context.user_data.get("pending_url", "")
            await query.message.reply_text("🌐 Fetching page...")
            text_content = fetch_url_content(url) or ""
            title = "Link: " + url[:60]
            if not text_content:
                text_content = f"(could not read page)\nURL: {url}"
            result, error = save_raw_to_drive(folder, title, f"URL: {url}\n\n{text_content}", src)
            await _finish_raw(query, context, folder, title, result, error)
            return

        if ptype == "file":
            # guard: Telegram won't let bots download files > 20 MB (get_file hangs)
            pending_size = context.user_data.get("pending_file_size", 0) or 0
            if pending_size > TELEGRAM_DOWNLOAD_LIMIT:
                size_mb = pending_size / (1024 * 1024)
                await query.message.reply_text(
                    f"⚠️ This file is {size_mb:.0f} MB. Telegram doesn't allow bots to download "
                    f"files larger than 20 MB, so I can't save it directly.\n\n"
                    f"Please compress it and resend:\n"
                    f"👉 https://www.ilovepdf.com/ru/compress_pdf\n\n"
                    f"(Or upload it to Google Drive manually and send me the link.)",
                    disable_web_page_preview=True
                )
                context.user_data.clear()
                return
            # download now, keep it, then ask for scores before saving
            try:
                await query.message.reply_text("📥 Downloading file...")
                tg_file = await asyncio.wait_for(
                    context.bot.get_file(context.user_data["pending_file_id"]),
                    timeout=45,
                )
                ext = os.path.splitext(fname)[1] or ".bin"
                fd, tmp_path = tempfile.mkstemp(suffix=ext)
                os.close(fd)
                await asyncio.wait_for(tg_file.download_to_drive(tmp_path), timeout=90)
            except asyncio.TimeoutError:
                await query.message.reply_text(
                    "⚠️ The download timed out. This usually means the file is too large for "
                    "Telegram to hand to a bot (the 20 MB limit).\n\n"
                    "Please compress it and resend:\n"
                    "👉 https://www.ilovepdf.com/ru/compress_pdf",
                    disable_web_page_preview=True
                )
                context.user_data.clear()
                return
            except Exception as e:
                await query.message.reply_text(
                    f"⚠️ Could not download the file: {str(e)[:200]}\n\n"
                    "If the file is larger than 20 MB, Telegram won't let me download it — "
                    "please compress it (ilovepdf.com/ru/compress_pdf) and resend."
                )
                context.user_data.clear()
                return

            context.user_data["raw_pending"] = {
                "folder": folder,
                "tmp_path": tmp_path,
                "fname": fname,
                "src": src,
            }
            await query.message.reply_text(
                f"📁 Folder: {folder}\n\n"
                "Before saving, set the scores. Send them as two numbers 1-10:\n"
                "importance credibility  — for example  7 8\n\n"
                "Or send 'skip' to save without scores."
            )
            return

    elif query.data.startswith("folder_"):
        context.user_data["pending_data"]["folder"] = query.data[7:]
        await query.edit_message_text(f"📁 Folder: {query.data[7:]}\n\nSend 'confirm' to save.")


async def complete_raw_file_save(update, context, importance, credibility):
    """Finish a raw file save once scores are provided (or skipped)."""
    rp = context.user_data.get("raw_pending")
    if not rp:
        return
    folder   = rp["folder"]
    tmp_path = rp["tmp_path"]
    fname    = rp["fname"]
    src      = rp["src"]

    stored_title = make_filename(fname or "Uploaded_file")
    try:
        size_mb = os.path.getsize(tmp_path) / (1024 * 1024)
    except Exception:
        size_mb = 0
    if size_mb > 10:
        await update.message.reply_text(
            f"💾 Saving to {folder}...\n"
            f"⏳ This is a large file ({size_mb:.0f} MB) — it may take 2–5 minutes. Please wait."
        )
    else:
        await update.message.reply_text(f"💾 Saving to {folder}...")
    drive_url, file_id, error = save_file_only_to_drive(
        folder, stored_title, tmp_path, fname, src,
        importance=importance, credibility=credibility
    )
    try: os.unlink(tmp_path)
    except Exception: pass

    if drive_url:
        context.bot_data.setdefault("docs", []).append({
            "title": stored_title, "folder": folder,
            "importance": importance or "—", "credibility": credibility or "—",
            "date": datetime.now().strftime("%Y-%m-%d"), "source": src,
            "drive_url": drive_url, "file_id": file_id or "", "original_id": "",
        })
        sc = f"⭐️{importance}/10  ✅{credibility}/10" if importance else "no scores"
        await update.message.reply_text(
            f"✅ File saved\n📁 {folder}\n📄 {stored_title}\n{sc}\n🔗 {drive_url}",
            disable_web_page_preview=True
        )
        await notify_admins(context, update,
            f"📢 {_display_name(update)} saved a file\n"
            f"📄 {stored_title}\n📁 {folder}\n🔗 {drive_url}")
    else:
        await update.message.reply_text(f"⚠️ Could not save: {error}")
    context.user_data.clear()


async def handle_score_edit(update, context):
    if not is_allowed(update):
        return
    # raw file waiting for scores?
    rp = context.user_data.get("raw_pending")
    if rp:
        text = update.message.text.strip().lower()
        if text.startswith("/skip") or text == "skip":
            await complete_raw_file_save(update, context, 0, 0)
            return
        m = re.match(r"(\d+)\s+(\d+)", text)
        if m:
            imp, cred = int(m.group(1)), int(m.group(2))
            if not (1 <= imp <= 10 and 1 <= cred <= 10):
                await update.message.reply_text("⚠️ Both scores must be 1–10. Try again, e.g. `7 8`.",
                                                parse_mode="Markdown")
                return
            await complete_raw_file_save(update, context, imp, cred)
        else:
            await update.message.reply_text(
                "Send two numbers 1–10: `importance credibility` (e.g. `7 8`), or /skip.",
                parse_mode="Markdown"
            )
        return
    text = update.message.text.strip().lower()
    pd   = context.user_data.get("pending_data")
    if not pd:
        return
    m = re.match(r"(importance|credibility)\s+(\d+)", text)
    if m:
        val = int(m.group(2))
        if 1 <= val <= 10:
            pd[m.group(1)] = val
            context.user_data["pending_data"] = pd
            await update.message.reply_text(
                f"✅ *{m.group(1).capitalize()}* → *{val}/10*\n\nSend `confirm` to save.",
                parse_mode="Markdown"
            )
        else:
            await update.message.reply_text("⚠️ Score must be 1–10.")
    elif text == "confirm":
        src = context.user_data.get("pending_source", "Unknown")
        pc  = context.user_data.get("pending_content", "")
        title  = pd.get("title", "Untitled")
        folder = pd.get("folder", "Knowledge/Market")
        orig_path = context.user_data.get("pending_original_path")
        orig_name = context.user_data.get("pending_original_name")
        original_url = ""
        original_id = ""
        if orig_path and os.path.exists(orig_path):
            await update.message.reply_text("💾 Saving original file and analysis...")
            drive_url, file_id, original_url, original_id, error = save_with_original_to_drive(
                pd, src, pc, orig_path, orig_name
            )
            try: os.unlink(orig_path)
            except Exception: pass
        else:
            await update.message.reply_text("💾 Saving to Google Drive...")
            drive_url, file_id, error = save_to_drive(pd, src, pc)
        context.bot_data.setdefault("docs", []).append({
            "title": title, "folder": folder,
            "importance": pd.get("importance", 5), "credibility": pd.get("credibility", 5),
            "date": pd.get("date", datetime.now().strftime("%Y-%m-%d")),
            "source": src, "drive_url": drive_url or "", "file_id": file_id or "",
            "original_id": original_id or "",
        })
        if drive_url:
            msg = f"✅ Saved!\n📁 {folder}\n📄 {title}\n🔗 {drive_url}"
            if original_url:
                msg += f"\n📎 Original: {original_url}"
            await update.message.reply_text(msg, disable_web_page_preview=True)
        else:
            await update.message.reply_text(f"⚠️ Drive error: {error}")
        context.user_data.clear()


async def setup_commands(app):
    """Register the command menu (the «/» button) and reset the Menu Button
    to show the command list — this overrides any stray BotFather/Web App
    menu button that may have been set externally."""
    await app.bot.set_my_commands([
        BotCommand("start",    "Welcome and overview"),
        BotCommand("help",     "How to use the bot"),
        BotCommand("status",   "Knowledge base stats"),
        BotCommand("list",     "Recent documents"),
        BotCommand("undo",     "Roll back the last save"),
        BotCommand("rollback", "Roll back a document by ID"),
        BotCommand("newfolder", "Create a folder (admins)"),
        BotCommand("delfolder", "Archive a folder (admins)"),
        BotCommand("toc", "Rebuild table of contents"),
    ])
    # force the blue Menu Button back to the default "commands" behaviour
    try:
        from telegram import MenuButtonCommands
        await app.bot.set_chat_menu_button(menu_button=MenuButtonCommands())
        logger.info("Menu button reset to commands")
    except Exception as e:
        logger.info(f"Could not reset menu button: {e}")
    logger.info("Bot command menu registered")


def main():
    app = Application.builder().token(TELEGRAM_TOKEN).post_init(setup_commands).build()

    # «Help ...» — a question to the knowledge base (not material to save)
    pomogi_filter = filters.TEXT & filters.Regex(r"(?i)^\s*help")
    # «Research: ...» — web research compiled into a document
    research_filter = filters.TEXT & filters.Regex(r"(?i)^\s*research[:\s]")

    # image documents (png/jpg sent as file) route to the photo handler
    image_doc_filter = filters.Document.MimeType("image/png") | filters.Document.MimeType("image/jpeg") | filters.Document.MimeType("image/webp")

    conv = ConversationHandler(
        entry_points=[
            MessageHandler(filters.TEXT & ~filters.COMMAND & ~pomogi_filter & ~research_filter, receive_message),
            MessageHandler(filters.PHOTO,                    handle_photo),
            MessageHandler(image_doc_filter,                 handle_photo),
            MessageHandler(filters.Document.ALL & ~image_doc_filter, receive_message),
            MessageHandler(filters.VOICE,                    handle_voice),
            MessageHandler(filters.AUDIO,                    handle_voice),
        ],
        states={WAITING_FOCUS: [MessageHandler(filters.TEXT & ~pomogi_filter & ~research_filter, receive_focus)]},
        fallbacks=[CommandHandler("skip", receive_focus)],
    )
    app.add_handler(CommandHandler("start",  start))
    app.add_handler(CommandHandler("help",   help_cmd))
    app.add_handler(CommandHandler("status", status))
    app.add_handler(CommandHandler("list",   list_docs))
    app.add_handler(CommandHandler("undo",     undo))
    app.add_handler(CommandHandler("rollback", rollback))
    app.add_handler(CommandHandler("newfolder", newfolder_cmd))
    app.add_handler(CommandHandler("delfolder", delfolder_cmd))
    app.add_handler(CommandHandler("setscore", setscore_cmd))
    app.add_handler(CommandHandler("toc", toc_cmd))
    app.add_handler(CommandHandler("whoisregistered", whoisregistered_cmd))
    app.add_handler(CommandHandler("testnotify", testnotify_cmd))
    app.add_handler(CommandHandler("drivetest", drivetest_cmd))
    # knowledge base question — before conv, to intercept first
    app.add_handler(MessageHandler(pomogi_filter, ask_knowledge_base))
    # web research — before conv
    app.add_handler(MessageHandler(research_filter, research))
    app.add_handler(conv)
    app.add_handler(CallbackQueryHandler(button_callback))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND & ~pomogi_filter & ~research_filter, handle_score_edit))
    logger.info("NeoBank Knowledge Bot (Phase 2 — Groq + Drive) starting...")
    app.run_polling(drop_pending_updates=True)

if __name__ == "__main__":
    main()
